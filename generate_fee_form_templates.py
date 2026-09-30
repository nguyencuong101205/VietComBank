# -*- coding: utf-8 -*-
import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

VCB_GREEN_HEX = "00843D"
VCB_LIGHT_GREEN_HEX = "E8F5E9"
BORDER_COLOR = "D0D7DE"

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def save_csv_and_excel(rows, csv_path, xlsx_path, sheet_name="Sheet1"):
    # 1. Ghi CSV
    os.makedirs(os.path.dirname(csv_path), exist_ok=True)
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerows(rows)

    # 2. Ghi Excel chuẩn màu thương hiệu Vietcombank
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=VCB_GREEN_HEX, end_color=VCB_GREEN_HEX, fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    for row_idx, row in enumerate(rows, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val_clean = str(val).strip()
            try:
                if '.' in val_clean:
                    cell.value = float(val_clean)
                else:
                    cell.value = int(val_clean)
            except ValueError:
                cell.value = val_clean

            cell.border = thin_border
            if row_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = data_font
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v = str(cell.value or '')
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.row_dimensions[1].height = 26
    os.makedirs(os.path.dirname(xlsx_path), exist_ok=True)
    wb.save(xlsx_path)

def create_word_doc(docx_path, title, form_code, subtitle, sections, table_data=None):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Header bảng
    tbl_h = doc.add_table(rows=1, cols=2)
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_h.columns[0].width = Inches(3.2)
    tbl_h.columns[1].width = Inches(3.5)

    c_left = tbl_h.cell(0, 0)
    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_bank = p_l.add_run("NGÂN HÀNG TMCP NGOẠI THƯƠNG VIỆT NAM\nVIETCOMBANK")
    r_bank.bold = True
    r_bank.font.size = Pt(10)
    r_bank.font.color.rgb = RGBColor(0, 132, 61)

    c_right = tbl_h.cell(0, 1)
    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_qh = p_r.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n---------------")
    r_qh.bold = True
    r_qh.font.size = Pt(10)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

    # Tiêu đề
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run(title.upper())
    r_t.bold = True
    r_t.font.size = Pt(14)
    r_t.font.color.rgb = RGBColor(0, 100, 45)

    if form_code:
        p_c = doc.add_paragraph()
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c = p_c.add_run(f"Mã văn bản / Biểu mẫu: {form_code}")
        r_c.italic = True
        r_c.font.size = Pt(9.5)
        r_c.font.color.rgb = RGBColor(100, 100, 100)

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.size = Pt(10.5)
        r_sub.bold = True

    # Nội dung các section
    for sec_title, fields in sections:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(3)
        r_sec = p_sec.add_run(sec_title)
        r_sec.bold = True
        r_sec.font.size = Pt(11)
        r_sec.font.color.rgb = RGBColor(0, 110, 50)

        for label, val in fields:
            p_f = doc.add_paragraph()
            p_f.paragraph_format.space_before = Pt(1)
            p_f.paragraph_format.space_after = Pt(2)
            r_lbl = p_f.add_run(f"• {label}: ")
            r_lbl.bold = True
            r_lbl.font.size = Pt(10)
            r_val = p_f.add_run(f"{val}")
            r_val.font.size = Pt(10)

    # Bảng chi tiết
    if table_data:
        p_tt = doc.add_paragraph()
        p_tt.paragraph_format.space_before = Pt(8)
        r_tt = p_tt.add_run("BẢNG KÊ QUY ĐỊNH CHI TIẾT:")
        r_tt.bold = True
        r_tt.font.size = Pt(10.5)

        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row_values in enumerate(table_data):
            for c_idx, val in enumerate(row_values):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(val)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.runs[0] if p.runs else p.add_run()
                run.font.size = Pt(9.5)
                if r_idx == 0:
                    run.bold = True
                    set_cell_background(cell, VCB_LIGHT_GREEN_HEX)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    if c_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Ký tên
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(14)

    tbl_sig = doc.add_table(rows=1, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.columns[0].width = Inches(3.3)
    tbl_sig.columns[1].width = Inches(3.4)

    c_s1 = tbl_sig.cell(0, 0)
    p_s1 = c_s1.paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s1 = p_s1.add_run("ĐẠI DIỆN KHÁCH HÀNG\n(Ký, ghi rõ họ tên & đóng dấu)\n\n\n\n...................................................")
    r_s1.font.size = Pt(10)
    r_s1.bold = True

    c_s2 = tbl_sig.cell(0, 1)
    p_s2 = c_s2.paragraphs[0]
    p_s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s2 = p_s2.add_run("ĐẠI DIỆN VIETCOMBANK\n(Giao dịch viên / Kiểm soát viên ký)\n\n\n\n...................................................")
    r_s2.font.size = Pt(10)
    r_s2.bold = True

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)

def main():
    # 5 tên chuẩn dựa trên ảnh chụp màn hình Biểu phí & biểu mẫu của hệ thống:
    # 1. Biểu phí dịch vụ Thẻ quốc tế Vietcombank Visa / MasterCard / JCB 2026
    # 2. Biểu mẫu đề nghị vay vốn & phương án trả nợ kiêm cam kết tài sản
    # 3. Bảng tổng hợp phí chuyển tiền quốc tế qua hệ thống SWIFT & Western Union
    # 4. Bảng biểu phí dịch vụ thẻ & tài khoản 2026
    # 5. Biểu mẫu đăng ký dịch vụ ngân hàng số VCB Digibank & Smart OTP 2026

    items = [
        {
            "id": "1",
            "base_name": "Bieu_phi_dich_vu_The_quoc_te_Vietcombank_Visa_MasterCard_JCB_2026",
            "title": "Biểu phí dịch vụ Thẻ quốc tế Vietcombank Visa / MasterCard / JCB 2026",
            "form_code": "VCB-CARD-INT-2026",
            "subtitle": "Quy định mức thu phí phát hành, phí thường niên và các phí giao dịch thẻ quốc tế",
            "table_data": [
                ["Hạng thẻ quốc tế", "Phí phát hành (VND)", "Phí thường niên (VND/năm)", "Phí chuyển đổi ngoại tệ", "Phí rút tiền ATM ngoài VN"],
                ["Vietcombank Visa Signature", "Miễn phí", "3,000,000", "1.50%", "4.0% (Tối thiểu 50,000 VND)"],
                ["Vietcombank Visa / MasterCard Platinum", "Miễn phí", "800,000", "2.00%", "4.0% (Tối thiểu 50,000 VND)"],
                ["Vietcombank JCB Platinum", "Miễn phí", "800,000", "2.00%", "4.0% (Tối thiểu 50,000 VND)"],
                ["Vietcombank Visa Hạng Vàng", "50,000", "200,000", "2.50%", "4.0% (Tối thiểu 50,000 VND)"],
                ["Vietcombank Visa Hạng Chuẩn", "50,000", "100,000", "2.50%", "4.0% (Tối thiểu 50,000 VND)"],
                ["Vietcombank Mastercard Chuẩn", "50,000", "100,000", "2.50%", "4.0% (Tối thiểu 50,000 VND)"],
                ["Vietcombank JCB Chuẩn", "50,000", "100,000", "2.50%", "4.0% (Tối thiểu 50,000 VND)"]
            ],
            "sections": [
                ("I. QUY ĐỊNH CHUNG VỀ THẺ QUỐC TẾ", [
                    ("Đối tượng áp dụng", "Khách hàng cá nhân và doanh nghiệp phát hành thẻ quốc tế tại Vietcombank"),
                    ("Thời hạn hiệu lực biểu phí", "Áp dụng từ ngày 01/01/2026 cho đến khi có thông báo thay thế mới"),
                    ("Quy định miễn/giảm phí", "Khách hàng VIP Vietcombank Priority được miễn 100% phí thường niên")
                ]),
                ("II. ĐẶC QUYỀN VÀ BẢO HIỂM DU LỊCH ĐI KÈM", [
                    ("Bảo hiểm du lịch toàn cầu", "Áp dụng cho chủ thẻ Visa Signature và Platinum hạn mức bồi thường tới 11.65 tỷ VND"),
                    ("Phòng chờ sân bay", "Miễn phí sử dụng phòng chờ Lotus Lounge tại các sân bay quốc tế Việt Nam"),
                    ("VCB Rewards", "Tích lũy điểm thưởng không giới hạn cho mọi giao dịch thanh toán POS/Online")
                ])
            ]
        },
        {
            "id": "2",
            "base_name": "Bieu_mau_de_nghi_vay_von_va_phuong_an_tra_no_kiem_cam_ket_tai_san",
            "title": "Biểu mẫu đề nghị vay vốn & phương án trả nợ kiêm cam kết tài sản",
            "form_code": "VCB-LOAN-PROP-2026",
            "subtitle": "Hồ sơ đăng ký vay tiêu dùng, vay mua bất động sản hoặc sản xuất kinh doanh",
            "table_data": [
                ["Kỳ thanh toán", "Dư nợ đầu kỳ (VND)", "Gốc phải trả (VND)", "Tiền lãi tạm tính (6.8%/năm)", "Tổng tiền trả kỳ (VND)"],
                ["Kỳ 1", "1,500,000,000", "12,500,000", "8,500,000", "21,000,000"],
                ["Kỳ 2", "1,487,500,000", "12,500,000", "8,429,167", "20,929,167"],
                ["Kỳ 3", "1,475,000,000", "12,500,000", "8,358,333", "20,858,333"],
                ["Kỳ 4", "1,462,500,000", "12,500,000", "8,287,500", "20,787,500"],
                ["Kỳ 5", "1,450,000,000", "12,500,000", "8,216,667", "20,716,667"],
                ["...", "...", "...", "...", "..."],
                ["Kỳ 120", "12,500,000", "12,500,000", "70,833", "12,570,833"]
            ],
            "sections": [
                ("I. THÔNG TIN KHÁCH HÀNG VÀ TÀI LIỆU PHÁP LÝ", [
                    ("Họ và tên khách hàng đề nghị vay", "NGUYỄN VĂN AN (Số CCCD: 001090012345)"),
                    ("Số điện thoại / Email", "0912345678 - an.nguyen@gmail.com"),
                    ("Địa chỉ thường trú / Nơi ở hiện tại", "Số 123 Phố Huế, Phường Hàng Bài, Quận Hoàn Kiếm, Hà Nội"),
                    ("Mục đích vay vốn", "Vay mua nhà ở dự án / căn hộ chung cư cao tầng")
                ]),
                ("II. THÔNG TIN KHOẢN VAY VÀ PHƯƠNG ÁN TRẢ NỢ", [
                    ("Số tiền đề nghị vay", "1,500,000,000 VND (Một tỷ năm trăm triệu đồng chẵn)"),
                    ("Thời hạn vay vốn", "120 tháng (10 năm)"),
                    ("Nguồn thu nhập trả nợ", "45,000,000 VND/tháng từ lương cố định sao kê Vietcombank"),
                    ("Tài sản bảo đảm thế chấp", "Căn hộ số 1204 Dự án Vinhomes Ocean Park (Giá trị thẩm định: 2.6 tỷ VND)")
                ])
            ]
        },
        {
            "id": "3",
            "base_name": "Bang_tong_hop_phi_chuyen_tien_quoc_te_qua_he_thong_SWIFT_va_Western_Union",
            "title": "Bảng tổng hợp phí chuyển tiền quốc tế qua hệ thống SWIFT & Western Union",
            "form_code": "VCB-REMIT-SWIFT-WU-2026",
            "subtitle": "Áp dụng cho dịch vụ chuyển tiền đi và nhận tiền từ nước ngoài qua SWIFT và Western Union",
            "table_data": [
                ["Dịch vụ chuyển tiền quốc tế", "Mức phí tiêu chuẩn", "Mức thu tối thiểu", "Mức thu tối đa", "Ghi chú phí ngân hàng đại lý"],
                ["SWIFT - Chuyển tiền đi (Phí SHA)", "0.20% số tiền chuyển", "5.00 USD", "200.00 USD", "Bên nhận chịu phí ngân hàng trung gian"],
                ["SWIFT - Chuyển tiền đi (Phí OUR)", "0.20% + 25 USD cước cố định", "30.00 USD", "250.00 USD", "Người gửi thanh toán toàn bộ chi phí"],
                ["SWIFT - Nhận tiền từ nước ngoài vào TK", "0.05% số tiền ghi có", "2.00 USD", "50.00 USD", "Miễn phí nếu nhận bằng VND"],
                ["Western Union - Nhận tiền quốc tế", "Miễn phí 100%", "0 USD", "0 USD", "Nhận tại quầy hoặc nhận qua VCB Digibank"],
                ["Western Union - Gửi tiền đi định cư/học tập", "Theo bậc tiền (0.3% - 0.5%)", "10.00 USD", "Theo biểu WU", "Nhanh chóng trong vòng vài phút"]
            ],
            "sections": [
                ("I. QUY ĐỊNH HỒ SƠ CHUYỂN TIỀN QUỐC TẾ HỢP PHÁP", [
                    ("Mục đích du học / Khám chữa bệnh", "Hộ chiếu, Visa, Thông báo học phí/viện phí hợp lệ"),
                    ("Mục đích trợ cấp thân nhân / Định cư", "Giấy tờ chứng minh quan hệ nhân thân và quyết định định cư"),
                    ("Hạn mức chuyển ngoại tệ hàng năm", "Theo quy định hướng dẫn của Ngân hàng Nhà nước Việt Nam")
                ]),
                ("II. THỜI GIAN VÀ TIẾN ĐỘ XỬ LÝ LỆNH", [
                    ("Điện chuyển tiền SWIFT GPI", "Theo dõi trạng thái lệnh chuyển tiền xuyên biên giới theo thời gian thực"),
                    ("Thời gian nhận tiền", "Từ 1 đến 2 ngày làm việc (với SWIFT) hoặc tức thì (với Western Union)")
                ])
            ]
        },
        {
            "id": "4",
            "base_name": "Bang_bieu_phi_dich_vu_the_va_tai_khoan_2026",
            "title": "Bảng biểu phí dịch vụ thẻ & tài khoản 2026",
            "form_code": "VCB-ACC-CARD-FEE-2026",
            "subtitle": "Áp dụng đối với Khách hàng Cá nhân sử dụng tài khoản thanh toán và thẻ nội địa Vietcombank",
            "table_data": [
                ["Dịch vụ tài khoản & thẻ", "Mức phí chuẩn (VND)", "Ưu đãi khách hàng Digibank", "Khách hàng Priority VIP", "Ghi chú"],
                ["Mở tài khoản thanh toán VND", "Miễn phí", "Miễn phí", "Miễn phí", "Mở tại quầy hoặc eKYC"],
                ["Duy trì số dư tài khoản tối thiểu", "0 VND", "0 VND", "0 VND", "Không yêu cầu số dư"],
                ["Phát hành thẻ ghi nợ Vietcombank Connect24", "50,000 VND", "Miễn phí khi mở online", "Miễn phí", "Thẻ chip Contactless"],
                ["Rút tiền tại ATM cùng hệ thống Vietcombank", "1,100 VND / GD", "1,100 VND / GD", "Miễn phí", "Đã bao gồm VAT"],
                ["Rút tiền tại ATM ngân hàng khác (Napas)", "3,300 VND / GD", "3,300 VND / GD", "Miễn phí", "Đã bao gồm VAT"],
                ["Chuyển tiền trong & ngoài hệ thống VCB", "0 VND", "0 VND", "0 VND", "Miễn phí trọn đời trên Digibank"],
                ["Quản lý tài khoản không hoạt động (>12 tháng)", "30,000 VND / tháng", "30,000 VND / tháng", "Miễn phí", "Chỉ thu khi tài khoản còn số dư"]
            ],
            "sections": [
                ("I. QUYỀN LỢI TÀI KHOẢN THANH TOÁN VIETCOMBANK", [
                    ("Hạn mức giao dịch chuyển tiền trực tuyến", "Lên đến 3 tỷ VND / ngày đối với khách hàng thông thường"),
                    ("Tài khoản số đẹp / Theo ngày sinh / SĐT", "Miễn phí chọn số tài khoản trùng số điện thoại trên VCB Digibank"),
                    ("Liên kết ví điện tử & Cổng thanh toán", "Miễn phí kết nối ShopeePay, MoMo, ZaloPay, Apple Pay, Google Pay")
                ]),
                ("II. AN TOÀN VÀ BẢO MẬT TÀI KHOẢN", [
                    ("Phương thức xác thực sinh trắc học", "Bắt buộc xác thực khuôn mặt khớp với chip CCCD cho các giao dịch trên 10 triệu VND"),
                    ("Tổng đài khóa thẻ / tài khoản khẩn cấp", "Hotline 1900 54 54 13 hoặc thao tác tự động 24/7 trên VCB Digibank")
                ])
            ]
        },
        {
            "id": "5",
            "base_name": "Bieu_mau_dang_ky_dich_vu_ngan_hang_so_VCB_Digibank_va_Smart_OTP_2026",
            "title": "Biểu mẫu đăng ký dịch vụ ngân hàng số VCB Digibank & Smart OTP 2026",
            "form_code": "VCB-EBANK-DIGI-2026",
            "subtitle": "Đăng ký mới, cấp đổi mã PIN, nâng hạn mức giao dịch và kích hoạt Smart OTP xác thực cao cấp",
            "table_data": [
                ["Gói dịch vụ Ngân hàng số", "Hạn mức tối đa / giao dịch (VND)", "Hạn mức tối đa / ngày (VND)", "Phương thức xác thực", "Phí duy trì"],
                ["Gói Cơ bản (Khởi đầu)", "100,000,000", "500,000,000", "SMS OTP", "Miễn phí"],
                ["Gói Nâng cao (Tiêu chuẩn)", "1,000,000,000", "3,000,000,000", "VCB Smart OTP tích hợp App", "Miễn phí"],
                ["Gói Priority VIP", "3,000,000,000", "10,000,000,000", "VCB Smart OTP + Sinh trắc học", "Miễn phí"],
                ["Gói Doanh nghiệp & Chủ hộ KD", "5,000,000,000", "20,000,000,000", "Hard Token OTP / Chữ ký số CA", "Miễn phí"]
            ],
            "sections": [
                ("I. THÔNG TIN KHÁCH HÀNG ĐĂNG KÝ", [
                    ("Họ và tên khách hàng", "LÊ HOÀNG MINH"),
                    ("Số CCCD gắn chip", "001201009876 (Ngày cấp: 15/09/2022)"),
                    ("Số điện thoại đăng ký Digibank", "0934567890"),
                    ("Địa chỉ hòm thư điện tử (Email)", "minh.lehoang@gmail.com")
                ]),
                ("II. DỊCH VỤ VÀ TÍNH NĂNG ĐĂNG KÝ MỚI", [
                    ("Loại yêu cầu", "[X] Đăng ký mới   [ ] Thay đổi thông tin SĐT   [ ] Mở khóa tài khoản"),
                    ("Gói hạn mức giao dịch đăng ký", "Gói Nâng cao (Hạn mức 3 tỷ VND / ngày)"),
                    ("Phương thức xác thực Smart OTP", "Kích hoạt Smart OTP trên thiết bị di động chính chủ"),
                    ("Dịch vụ nhận thông báo OTT", "Nhận thông báo biến động số dư miễn phí tức thì trên VCB Digibank")
                ]),
                ("III. CAM KẾT VÀ BẢO MẬT CỦA KHÁCH HÀNG", [
                    ("Bảo mật mật khẩu và mã OTP", "Tuyệt đối không cung cấp mật khẩu, mã OTP, mã PIN cho bất kỳ ai kể cả nhân viên ngân hàng"),
                    ("Cập nhật thông tin sinh trắc học", "Khách hàng đã hoàn thành quét chip CCCD thành công trên ứng dụng VCB Digibank")
                ])
            ]
        }
    ]

    target_main_dir = r"c:\Users\Admin\Documents\Đồ_Án\BTL\Du_Lieu_Mau_He_Thong"
    public_samples_dir = r"c:\Users\Admin\Documents\Đồ_Án\BTL\frontend\public\samples"
    samples_dir = r"c:\Users\Admin\Documents\Đồ_Án\BTL\samples"

    dirs_to_write = [
        target_main_dir,
        public_samples_dir,
        samples_dir
    ]

    # Tạo các thư mục con phân loại trong Du_Lieu_Mau_He_Thong
    excel_dir = os.path.join(target_main_dir, "1_File_Excel_XLSX")
    csv_dir = os.path.join(target_main_dir, "2_File_CSV")
    word_dir = os.path.join(target_main_dir, "3_File_Word_DOCX")
    os.makedirs(excel_dir, exist_ok=True)
    os.makedirs(csv_dir, exist_ok=True)
    os.makedirs(word_dir, exist_ok=True)

    for it in items:
        base = it["base_name"]
        table = it["table_data"]
        title = it["title"]
        code = it["form_code"]
        sub = it["subtitle"]
        secs = it["sections"]

        # 1. Tạo file Excel (.xlsx) & CSV (.csv)
        csv_file = f"{base}.csv"
        xlsx_file = f"{base}.xlsx"
        docx_file = f"{base}.docx"

        # Ghi vào thư mục phân loại riêng
        save_csv_and_excel(table, os.path.join(csv_dir, csv_file), os.path.join(excel_dir, xlsx_file), sheet_name="Data")
        create_word_doc(os.path.join(word_dir, docx_file), title, code, sub, secs, table)

        # Ghi vào thư mục tổng hợp Du_Lieu_Mau_He_Thong và public samples
        for d in dirs_to_write:
            save_csv_and_excel(table, os.path.join(d, csv_file), os.path.join(d, xlsx_file), sheet_name="Data")
            create_word_doc(os.path.join(d, docx_file), title, code, sub, secs, table)

        print(f"-> Hoan tat: {base} (Excel, CSV, Word)")

    print("\nSUCCESS_ALL_TEMPLATES_CREATED")

if __name__ == "__main__":
    main()
