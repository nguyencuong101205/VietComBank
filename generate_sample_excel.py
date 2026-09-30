import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def csv_to_excel(csv_path, xlsx_path, sheet_name="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Styling definitions (Vietcombank green theme)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="00843D", end_color="00843D", fill_type="solid") # VCB Green
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    with open(csv_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, 1):
            for col_idx, val in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx)

                # Attempt to convert numeric types
                val_clean = val.strip()
                try:
                    if '.' in val_clean:
                        cell.value = float(val_clean)
                    else:
                        cell.value = int(val_clean)
                except ValueError:
                    cell.value = val_clean

                cell.border = thin_border

                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = data_font
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.row_dimensions[1].height = 25
    wb.save(xlsx_path)
    print(f"Generated {xlsx_path}")

def main():
    source_dir = "samples"
    public_dir = "frontend/public/samples"
    os.makedirs(public_dir, exist_ok=True)

    files = [
        ("sample_exchange_rates.csv", "sample_exchange_rates.xlsx", "ExchangeRates"),
        ("sample_gold_rates.csv", "sample_gold_rates.xlsx", "GoldRates"),
        ("sample_interest_rates.csv", "sample_interest_rates.xlsx", "InterestRates"),
        ("sample_payroll_employees.csv", "sample_payroll_employees.xlsx", "PayrollEmployees"),
    ]

    for csv_file, xlsx_file, sheet_name in files:
        csv_path = os.path.join(source_dir, csv_file)
        xlsx_path_samples = os.path.join(source_dir, xlsx_file)
        xlsx_path_public = os.path.join(public_dir, xlsx_file)

        csv_to_excel(csv_path, xlsx_path_samples, sheet_name)
        csv_to_excel(csv_path, xlsx_path_public, sheet_name)

if __name__ == "__main__":
    main()
